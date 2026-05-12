from __future__ import annotations

import io
from pathlib import Path

from fastapi import HTTPException, UploadFile

from .schemas import ALLOWED_EXTENSIONS, ParsedFile


async def parse_upload(file: UploadFile) -> ParsedFile:
    ext = Path(file.filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=422,
            detail=(
                f"File '{file.filename}': extension '{ext}' is not allowed. "
                f"Allowed: {sorted(ALLOWED_EXTENSIONS)}"
            ),
        )
    data = await file.read()
    text = _extract(ext, data)
    return ParsedFile(filename=file.filename or "unnamed", text=text)


def _extract(ext: str, data: bytes) -> str:
    if ext == ".pdf":
        return _pdf(data)
    if ext in {".doc", ".docx"}:
        return _docx(data)
    if ext in {".xls", ".xlsx"}:
        return _excel(data)
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)


def _excel(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            lines.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)
